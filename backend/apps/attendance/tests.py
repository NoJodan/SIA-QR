"""Tests de regresión QA v2 (C2, C3, M1, M2, M3)."""

import hashlib
import io
import uuid
from datetime import timedelta
from unittest.mock import patch

from django.test import TestCase
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.academic.models import (
    AcademicGroup,
    ClassStatus,
    Course,
    ScheduledClass,
)
from apps.attendance.models import Attendance, AttendanceSession, AuditLog
from apps.authentication.models import (
    AuthorizedProfessorEmail,
    Professor,
    Student,
    User,
)


def _make_prof_group(email="prof@ut.edu.co"):
    user = User.objects.create_user(email=email, google_sub=f"sub-{email}", role="ROLE_PROFESSOR")
    prof = Professor.objects.create(
        user=user, employee_code=f"EMP-{email[:3].upper()}", first_name="Ana", last_name="Perez"
    )
    course = Course.objects.create(code=f"C-{email[:3].upper()}", name="Curso QA")
    group = AcademicGroup.objects.create(
        course=course, professor=prof, group_code="01", term_period="2026-1"
    )
    return user, prof, group


def _make_class(group, status=ClassStatus.IN_PROGRESS):
    return ScheduledClass.objects.create(
        group=group,
        title="Clase QA",
        start_time=timezone.now(),
        duration_minutes=60,
        qr_duration_minutes=10,
        modality="PRESENTIAL",
        status=status,
    )


def _make_session(clase, active=True):
    raw = f"tok-{clase.id}-{AttendanceSession.objects.count()}"
    return AttendanceSession.objects.create(
        scheduled_class=clase,
        token_hash=hashlib.sha256(raw.encode()).hexdigest(),
        expires_at=timezone.now() + timedelta(minutes=10),
        is_active=active,
    ), raw


def _make_student(email, code="EST-QA", doc="DOC-QA"):
    user = User.objects.create_user(
        email=email, google_sub=f"sub-{email}", role="ROLE_STUDENT"
    )
    student = Student.objects.create(
        user=user, student_code=code, document_number=doc,
        first_name="Est", last_name="QA",
        phone_number="+57 300 123 4567", address="Calle 1 #2-3",
    )
    return user, student


class MarkGeoToleranceTests(TestCase):
    """RNF-06: la geo nunca bloquea la marcación (no 400 por decimales)."""

    def test_serializer_rounds_15_decimals(self):
        from apps.attendance.serializers import MarkAttendanceSerializer

        ser = MarkAttendanceSerializer(data={
            "token": "tok",
            "latitude": 4.123456789012345,
            "longitude": -75.123456789012345,
            "accuracy": 12.3456789,
        })
        self.assertTrue(ser.is_valid(), ser.errors)
        self.assertEqual(str(ser.validated_data["latitude"]), "4.12345679")
        self.assertEqual(str(ser.validated_data["longitude"]), "-75.12345679")
        self.assertEqual(str(ser.validated_data["accuracy"]), "12.35")

    def test_serializer_accepts_str_geo(self):
        from apps.attendance.serializers import MarkAttendanceSerializer

        ser = MarkAttendanceSerializer(data={
            "token": "tok",
            "latitude": "4.123456789012345",
            "longitude": "-75.123456789012345",
            "accuracy": "3.14159",
        })
        self.assertTrue(ser.is_valid(), ser.errors)
        self.assertEqual(str(ser.validated_data["latitude"]), "4.12345679")
        self.assertEqual(str(ser.validated_data["accuracy"]), "3.14")

    def test_serializer_null_empty_missing_pass(self):
        from apps.attendance.serializers import MarkAttendanceSerializer

        for geo in (
            {"latitude": None, "longitude": None, "accuracy": None},
            {"latitude": "", "longitude": "", "accuracy": ""},
            {},
        ):
            ser = MarkAttendanceSerializer(data={"token": "tok", **geo})
            self.assertTrue(ser.is_valid(), ser.errors)
            self.assertIsNone(ser.validated_data.get("latitude"))
            self.assertIsNone(ser.validated_data.get("longitude"))
            self.assertIsNone(ser.validated_data.get("accuracy"))

    def test_serializer_out_of_range_coerces_to_none(self):
        from apps.attendance.serializers import MarkAttendanceSerializer

        ser = MarkAttendanceSerializer(data={
            "token": "tok", "latitude": 91, "longitude": 200, "accuracy": -5,
        })
        self.assertTrue(ser.is_valid(), ser.errors)
        self.assertIsNone(ser.validated_data.get("latitude"))
        self.assertIsNone(ser.validated_data.get("longitude"))
        self.assertIsNone(ser.validated_data.get("accuracy"))

    def test_serializer_garbage_coerces_to_none(self):
        from apps.attendance.serializers import MarkAttendanceSerializer

        ser = MarkAttendanceSerializer(data={
            "token": "tok", "latitude": "NaN", "longitude": "abc", "accuracy": "Infinity",
        })
        self.assertTrue(ser.is_valid(), ser.errors)
        self.assertIsNone(ser.validated_data.get("latitude"))
        self.assertIsNone(ser.validated_data.get("longitude"))
        self.assertIsNone(ser.validated_data.get("accuracy"))

    def test_serializer_token_still_strict(self):
        from apps.attendance.serializers import MarkAttendanceSerializer

        for bad in ("", "   "):
            ser = MarkAttendanceSerializer(data={"token": bad})
            self.assertFalse(ser.is_valid())
            self.assertIn("token", ser.errors)
        ser = MarkAttendanceSerializer(data={})
        self.assertFalse(ser.is_valid())
        self.assertIn("token", ser.errors)


class MarkAttendanceGeoMatrixTests(TestCase):
    """Matriz POST /api/attendance/mark/ con/sin geo y códigos de estado."""

    def _mark(self, email, raw, payload_extra=None, code="EST-M", doc="DOC-M"):
        user, _ = _make_student(email, code=code, doc=doc)
        client = APIClient()
        client.force_login(user)
        body = {"token": raw}
        if payload_extra:
            body.update(payload_extra)
        return client.post("/api/attendance/mark/", body, format="json")

    def test_mark_with_precise_geo_is_201(self):
        _, _, group = _make_prof_group(email="geoprof@ut.edu.co")
        _, raw = _make_session(_make_class(group))
        resp = self._mark("geo1@ut.edu.co", raw, {
            "latitude": 4.123456789012345,
            "longitude": -75.1234567890123,
            "accuracy": 12.3456789,
        }, code="EST-G1", doc="DOC-G1")
        self.assertEqual(resp.status_code, 201)
        self.assertFalse(resp.data.get("already_marked"))

    def test_mark_without_geo_is_201(self):
        _, _, group = _make_prof_group(email="nogeo@ut.edu.co")
        _, raw = _make_session(_make_class(group))
        resp = self._mark("geo2@ut.edu.co", raw,
                          {"latitude": None, "longitude": None, "accuracy": None},
                          code="EST-G2", doc="DOC-G2")
        self.assertEqual(resp.status_code, 201)

    def test_mark_with_long_accuracy_is_201(self):
        _, _, group = _make_prof_group(email="accprof@ut.edu.co")
        _, raw = _make_session(_make_class(group))
        resp = self._mark("geo3@ut.edu.co", raw, {
            "latitude": 4.5, "longitude": -75.5, "accuracy": 9.87654321,
        }, code="EST-G3", doc="DOC-G3")
        self.assertEqual(resp.status_code, 201)

    def test_mark_invalid_token_is_404(self):
        resp = self._mark("geo4@ut.edu.co", "token-inexistente",
                          code="EST-G4", doc="DOC-G4")
        # Sin sesión existente el flujo llega a token->404 (el perfil existe).
        self.assertEqual(resp.status_code, 404)

    def test_mark_expired_is_410(self):
        _, _, group = _make_prof_group(email="expprof@ut.edu.co")
        session, raw = _make_session(_make_class(group))
        session.expires_at = timezone.now() - timedelta(minutes=1)
        session.save(update_fields=["expires_at"])
        resp = self._mark("geo5@ut.edu.co", raw, code="EST-G5", doc="DOC-G5")
        self.assertEqual(resp.status_code, 410)

    def test_mark_without_profile_is_412(self):
        _, _, group = _make_prof_group(email="noprof@ut.edu.co")
        _, raw = _make_session(_make_class(group))
        user = User.objects.create_user(
            email="sinperfil@ut.edu.co", google_sub="sub-sinperfil",
            role="ROLE_STUDENT",
        )
        client = APIClient()
        client.force_login(user)
        resp = client.post("/api/attendance/mark/", {"token": raw}, format="json")
        self.assertEqual(resp.status_code, 412)

    def test_mark_duplicate_is_200_already_marked(self):
        _, _, group = _make_prof_group(email="dupprof@ut.edu.co")
        _, raw = _make_session(_make_class(group))
        user, _ = _make_student("geo6@ut.edu.co", code="EST-G6", doc="DOC-G6")
        client = APIClient()
        client.force_login(user)
        first = client.post("/api/attendance/mark/",
                            {"token": raw, "latitude": 4.123456789012345,
                             "longitude": -75.1234567890123, "accuracy": 5.55555},
                            format="json")
        self.assertEqual(first.status_code, 201)
        second = client.post("/api/attendance/mark/", {"token": raw}, format="json")
        self.assertEqual(second.status_code, 200)
        self.assertTrue(second.data.get("already_marked"))

    def test_mark_closed_class_is_409(self):
        _, _, group = _make_prof_group(email="closedprof@ut.edu.co")
        _, raw = _make_session(_make_class(group, status=ClassStatus.COMPLETED))
        resp = self._mark("geo7@ut.edu.co", raw, code="EST-G7", doc="DOC-G7")
        self.assertEqual(resp.status_code, 409)


class SingleActiveConstraintTests(TestCase):
    def test_two_active_sessions_same_class_violates_constraint(self):
        _, _, group = _make_prof_group()
        clase = _make_class(group)
        _make_session(clase)
        from django.db import IntegrityError

        with self.assertRaises(IntegrityError):
            _make_session(clase)

    def test_generate_session_keeps_single_active(self):
        from apps.attendance.services import generate_session

        _, _, group = _make_prof_group(email="prof2@ut.edu.co")
        clase = _make_class(group, status=ClassStatus.SCHEDULED)
        generate_session(clase)
        generate_session(clase)
        self.assertEqual(
            AttendanceSession.objects.filter(scheduled_class=clase, is_active=True).count(), 1
        )


class ClosedClassTests(TestCase):
    def test_rotate_on_completed_class_is_410(self):
        user, _, group = _make_prof_group(email="prof3@ut.edu.co")
        clase = _make_class(group, status=ClassStatus.COMPLETED)
        client = APIClient()
        client.force_login(user)
        resp = client.post(
            f"/api/academic/groups/{group.id}/classes/{clase.id}/current-session/rotate/",
            {},
            format="json",
        )
        self.assertEqual(resp.status_code, 410)
        self.assertEqual(resp.data["state"], "finished")

    def test_mark_on_cancelled_class_is_409(self):
        _, _, group = _make_prof_group(email="prof4@ut.edu.co")
        clase = _make_class(group, status=ClassStatus.CANCELLED)
        session, raw = _make_session(clase)
        student_user = User.objects.create_user(
            email="stu@ut.edu.co", google_sub="sub-stu", role="ROLE_STUDENT"
        )
        Student.objects.create(
            user=student_user, student_code="EST-1", document_number="DOC-1",
            first_name="Est", last_name="Uno",
        )
        client = APIClient()
        client.force_login(student_user)
        resp = client.post("/api/attendance/mark/", {"token": raw}, format="json")
        self.assertEqual(resp.status_code, 409)

    def test_resolve_on_completed_class_is_410(self):
        _, _, group = _make_prof_group(email="prof5@ut.edu.co")
        clase = _make_class(group, status=ClassStatus.COMPLETED)
        _, raw = _make_session(clase)
        student_user = User.objects.create_user(
            email="stu2@ut.edu.co", google_sub="sub-stu2", role="ROLE_STUDENT"
        )
        client = APIClient()
        client.force_login(student_user)
        resp = client.get("/api/attendance/resolve/", {"token": raw})
        self.assertEqual(resp.status_code, 410)


class ClassWindowEnforcementTests(TestCase):
    """El TTL largo del QR no permite marcar después del fin de la clase."""

    def _setup(self, email):
        _, _, group = _make_prof_group(email=email)
        clase = ScheduledClass.objects.create(
            group=group,
            title="Ventana vencida",
            start_time=timezone.now() - timedelta(minutes=10),
            duration_minutes=5,
            qr_duration_minutes=120,
            modality="PRESENTIAL",
            status=ClassStatus.IN_PROGRESS,
        )
        _, raw = _make_session(clase)
        return clase, raw

    def test_mark_after_end_is_409(self):
        _, raw = self._setup("w1@ut.edu.co")
        user, _ = _make_student("w1est@ut.edu.co", code="EST-W1", doc="DOC-W1")
        client = APIClient()
        client.force_login(user)
        resp = client.post("/api/attendance/mark/", {"token": raw}, format="json")
        self.assertEqual(resp.status_code, 409)

    def test_resolve_after_end_is_410_finished(self):
        clase, raw = self._setup("w2@ut.edu.co")
        user, _ = _make_student("w2est@ut.edu.co", code="EST-W2", doc="DOC-W2")
        client = APIClient()
        client.force_login(user)
        resp = client.get("/api/attendance/resolve/", {"token": raw})
        self.assertEqual(resp.status_code, 410)
        self.assertEqual(resp.data["state"], "finished")
        self.assertEqual(resp.data["reason"], "window_elapsed")
        # Un GET no cierra la clase (cierre perezoso solo en el poll).
        clase.refresh_from_db()
        self.assertEqual(clase.status, ClassStatus.IN_PROGRESS)


class StudentProfileTests(TestCase):
    def test_mark_without_profile_is_412_needs_profile(self):
        _, _, group = _make_prof_group(email="prof6@ut.edu.co")
        clase = _make_class(group)
        _, raw = _make_session(clase)
        student_user = User.objects.create_user(
            email="noprofile@ut.edu.co", google_sub="sub-np", role="ROLE_STUDENT"
        )
        client = APIClient()
        client.force_login(student_user)
        resp = client.post("/api/attendance/mark/", {"token": raw}, format="json")
        self.assertEqual(resp.status_code, 412)
        self.assertTrue(resp.data.get("needs_profile"))

    def test_patch_student_profile_creates_real_profile(self):
        from allauth.socialaccount.models import SocialAccount

        student_user = User.objects.create_user(
            email="newstu@ut.edu.co", google_sub="sub-new", role="ROLE_STUDENT"
        )
        SocialAccount.objects.create(
            user=student_user, provider="google", uid="sub-new",
            extra_data={"given_name": "Real", "family_name": "Estudiante"},
        )
        client = APIClient()
        client.force_login(student_user)
        resp = client.patch(
            "/api/auth/me/",
            {"student_code": "123456789012", "document_number": "123456",
             "phone_number": "3001234567", "address": "Calle 1 #2-3",
             "first_name": "Real", "last_name": "Estudiante"},
            format="json",
        )
        self.assertEqual(resp.status_code, 200)
        self.assertFalse(resp.data.get("needs_profile"))
        self.assertEqual(Student.objects.get(user=student_user).document_number, "123456")

    def test_patch_student_duplicate_is_409(self):
        from allauth.socialaccount.models import SocialAccount

        u1 = User.objects.create_user(email="s1@ut.edu.co", google_sub="s1", role="ROLE_STUDENT")
        Student.objects.create(
            user=u1, student_code="123456789012", document_number="123456789012",
            first_name="A", last_name="B", phone_number="3001111111",
            address="Dir 1",
        )
        u2 = User.objects.create_user(email="s2@ut.edu.co", google_sub="s2", role="ROLE_STUDENT")
        SocialAccount.objects.create(
            user=u2, provider="google", uid="s2",
            extra_data={"given_name": "C", "family_name": "D"},
        )
        client = APIClient()
        client.force_login(u2)
        resp = client.patch(
            "/api/auth/me/",
            {"student_code": "123456789012", "document_number": "123456789013",
             "phone_number": "3002222222", "address": "Dir 2",
             "first_name": "C", "last_name": "D"},
            format="json",
        )
        self.assertEqual(resp.status_code, 409)


class ProfessorReconcileTests(TestCase):
    def test_me_promotes_late_whitelisted_student(self):
        user = User.objects.create_user(
            email="late@ut.edu.co", google_sub="sub-late", role="ROLE_STUDENT"
        )
        AuthorizedProfessorEmail.objects.create(email="late@ut.edu.co")
        client = APIClient()
        client.force_login(user)
        resp = client.get("/api/auth/me/")
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.data["role"], "ROLE_PROFESSOR")
        user.refresh_from_db()
        self.assertEqual(user.role, "ROLE_PROFESSOR")
        self.assertTrue(Professor.objects.filter(user=user).exists())


# Referencia GPS del salón (4.0915, -75.2315); las coordenadas de los alumnos
# se eligen respecto a ella para los casos dentro/fuera del rango de 25 m.
SALON_LAT, SALON_LNG = 4.0915, -75.2315


def _set_geofence(clase, lat=SALON_LAT, lng=SALON_LNG, radius=25):
    clase.geofence_lat = lat
    clase.geofence_lng = lng
    clase.geofence_radius_m = radius
    clase.geofence_activated_at = timezone.now()
    clase.save(update_fields=[
        "geofence_lat", "geofence_lng", "geofence_radius_m", "geofence_activated_at", "updated_at",
    ])
    return clase


class HaversineTests(TestCase):
    def test_same_point_is_zero(self):
        from apps.attendance.services import haversine_m

        self.assertAlmostEqual(haversine_m(SALON_LAT, SALON_LNG, SALON_LAT, SALON_LNG), 0.0, places=3)

    def test_known_distance_north(self):
        from apps.attendance.services import haversine_m

        # ~111.2 m al norte (0.001 grado de latitud).
        distance = haversine_m(SALON_LAT, SALON_LNG, SALON_LAT + 0.001, SALON_LNG)
        self.assertAlmostEqual(distance, 111.2, delta=1.0)

    def test_none_coordinates_return_none(self):
        from apps.attendance.services import haversine_m

        self.assertIsNone(haversine_m(None, SALON_LNG, SALON_LAT, SALON_LNG))
        self.assertIsNone(haversine_m(SALON_LAT, SALON_LNG, None, None))

    def test_antipodal_coordinates_do_not_crash(self):
        from apps.attendance.services import haversine_m

        # a redondea a 1.0000000000000002: sin el clamp, sqrt(1-a) revienta.
        distance = haversine_m(45, 2.154007, -45, -177.845993)
        self.assertAlmostEqual(distance, 20015114.44, delta=1.0)


class AttendanceRangeTests(TestCase):
    """inside_range/distance_m de la lista viva (profesor dueño)."""

    def _setup(self, email="rangeprof@ut.edu.co", radius=25, geofence=True, modality="PRESENTIAL"):
        owner, _, group = _make_prof_group(email=email)
        clase = _make_class(group, status=ClassStatus.IN_PROGRESS)
        clase.modality = modality
        clase.save(update_fields=["modality"])
        if geofence:
            _set_geofence(clase, radius=radius)
        session, raw = _make_session(clase)
        client = APIClient()
        client.force_login(owner)
        return client, session, raw, clase, group

    def _mark(self, raw, email, lat=None, lng=None):
        user, _ = _make_student(
            email, code=f"EST-{email[:4].upper()}", doc=f"DOC-{email[:4].upper()}"
        )
        client = APIClient()
        client.force_login(user)
        body = {"token": raw}
        if lat is not None:
            body["latitude"] = lat
            body["longitude"] = lng
            body["accuracy"] = 8.0
        return client.post("/api/attendance/mark/", body, format="json")

    def _add_attendance(self, session, student, lat, lng):
        from apps.attendance.models import Attendance

        return Attendance.objects.create(
            session=session, student=student,
            latitude=lat, longitude=lng, accuracy=5,
            ip_address="127.0.0.1", user_agent="test",
        )

    def _rows(self, owner_client, session):
        resp = owner_client.get(f"/api/attendance/sessions/{session.id}/attendances/")
        self.assertEqual(resp.status_code, 200)
        return resp.data["results"] if isinstance(resp.data, dict) else resp.data

    def test_inside_and_outside_range(self):
        client, session, _, _, _ = self._setup(email="r1prof@ut.edu.co")
        _, student_in = _make_student("dentro@ut.edu.co", code="EST-IN", doc="DOC-IN")
        _, student_out = _make_student("fuera@ut.edu.co", code="EST-OUT", doc="DOC-OUT")
        self._add_attendance(session, student_in, SALON_LAT, SALON_LNG)
        self._add_attendance(session, student_out, SALON_LAT + 0.001, SALON_LNG)
        rows = self._rows(client, session)
        by_code = {r["student_code"]: r for r in rows}
        self.assertTrue(by_code["EST-IN"]["inside_range"])
        self.assertAlmostEqual(by_code["EST-IN"]["distance_m"], 0.0, places=1)
        self.assertFalse(by_code["EST-OUT"]["inside_range"])
        self.assertGreater(by_code["EST-OUT"]["distance_m"], 25)

    def test_outside_radius_still_marks_attendance(self):
        # RNF-06: fuera del rango NO bloquea; solo se etiqueta.
        client, session, raw, _, _ = self._setup(email="r2prof@ut.edu.co")
        resp = self._mark(raw, "fuera2@ut.edu.co", lat=SALON_LAT + 0.001, lng=SALON_LNG)
        self.assertEqual(resp.status_code, 201)
        rows = self._rows(client, session)
        self.assertEqual(len(rows), 1)
        self.assertFalse(rows[0]["inside_range"])

    def test_no_location_with_range_is_outside(self):
        # Rango fijado y alumno sin GPS: cuenta como Fuera de Rango.
        client, session, raw, _, _ = self._setup(email="r3prof@ut.edu.co")
        resp = self._mark(raw, "sinloc@ut.edu.co")
        self.assertEqual(resp.status_code, 201)  # la marcación nunca se bloquea
        rows = self._rows(client, session)
        self.assertEqual(len(rows), 1)
        self.assertFalse(rows[0]["inside_range"])
        self.assertIsNone(rows[0]["distance_m"])
        self.assertFalse(rows[0]["has_location"])

    def test_null_when_class_has_no_geofence(self):
        client, session, raw, _, _ = self._setup(email="r4prof@ut.edu.co", geofence=False)
        self._mark(raw, "sinrango@ut.edu.co", lat=SALON_LAT, lng=SALON_LNG)
        rows = self._rows(client, session)
        self.assertIsNone(rows[0]["inside_range"])
        self.assertIsNone(rows[0]["distance_m"])
        # Sigue indicando que sí envió ubicación.
        self.assertTrue(rows[0]["has_location"])

    def test_null_when_class_is_virtual(self):
        client, session, raw, _, _ = self._setup(email="r5prof@ut.edu.co", modality="VIRTUAL")
        self._mark(raw, "mediada@ut.edu.co", lat=SALON_LAT, lng=SALON_LNG)
        rows = self._rows(client, session)
        self.assertIsNone(rows[0]["inside_range"])

    def test_verdict_survives_after_close(self):
        # El veredicto se re-valida en cada lectura: sigue visible al cerrar
        # la clase (la geocerca ya no cambia después del cierre).
        client, session, _, clase, _ = self._setup(email="r6prof@ut.edu.co")
        _, student_in = _make_student("cerro-in@ut.edu.co", code="EST-CI", doc="DOC-CI")
        _, student_out = _make_student("cerro-out@ut.edu.co", code="EST-CO", doc="DOC-CO")
        self._add_attendance(session, student_in, SALON_LAT, SALON_LNG)
        self._add_attendance(session, student_out, SALON_LAT + 0.001, SALON_LNG)
        clase.status = ClassStatus.COMPLETED
        clase.save(update_fields=["status", "updated_at"])
        rows = {r["student_code"]: r for r in self._rows(client, session)}
        self.assertTrue(rows["EST-CI"]["inside_range"])
        self.assertAlmostEqual(rows["EST-CI"]["distance_m"], 0.0, places=1)
        self.assertFalse(rows["EST-CO"]["inside_range"])
        self.assertGreater(rows["EST-CO"]["distance_m"], 25)

    def test_custom_radius_respected(self):
        client, session, _, _, _ = self._setup(email="r7prof@ut.edu.co", radius=200)
        _, student = _make_student("lejos@ut.edu.co", code="EST-LEJ", doc="DOC-LEJ")
        self._add_attendance(session, student, SALON_LAT + 0.001, SALON_LNG)
        rows = self._rows(client, session)
        # ~111 m: fuera de 25 m pero dentro de un radio ampliado a 200 m.
        self.assertTrue(rows[0]["inside_range"])

    def test_exact_coords_do_not_expose_lat_long(self):
        client, session, _, _, _ = self._setup(email="r8prof@ut.edu.co")
        _, student = _make_student("priv@ut.edu.co", code="EST-PRIV", doc="DOC-PRIV")
        self._add_attendance(session, student, SALON_LAT, SALON_LNG)
        rows = self._rows(client, session)
        self.assertNotIn("latitude", rows[0])
        self.assertNotIn("longitude", rows[0])


class AuditLogTests(TestCase):
    """RF-ADM-07: eventos de auditoría de marcaciones."""

    def _mark(self, email, code, doc):
        _, _, group = _make_prof_group(email=f"audit{email}")
        session, raw = _make_session(_make_class(group))
        user, _ = _make_student(email, code=code, doc=doc)
        client = APIClient()
        client.force_login(user)
        resp = client.post(
            "/api/attendance/mark/", {"token": raw}, format="json"
        )
        return user, session, resp

    def test_mark_creates_audit_log(self):
        user, session, resp = self._mark(
            "a1@ut.edu.co", "EST-AUD1", "DOC-AUD1"
        )
        self.assertEqual(resp.status_code, 201)
        log = AuditLog.objects.get(event_type="attendance_marked")
        self.assertEqual(log.actor_id, user.id)
        self.assertEqual(log.resource, f"session:{session.id}")
        self.assertNotIn("latitude", log.payload)
        self.assertNotIn("longitude", log.payload)

    def test_duplicate_mark_logs_duplicate_event(self):
        _, _, group = _make_prof_group(email="auditdup@ut.edu.co")
        session, raw = _make_session(_make_class(group))
        user, _ = _make_student("a2@ut.edu.co", code="EST-AUD2", doc="DOC-AUD2")
        client = APIClient()
        client.force_login(user)
        client.post("/api/attendance/mark/", {"token": raw}, format="json")
        client.post("/api/attendance/mark/", {"token": raw}, format="json")
        event = AuditLog.objects.get(event_type="attendance_duplicate")
        self.assertEqual(event.resource, f"session:{session.id}")

    @patch(
        "apps.attendance.models.AuditLog.objects.create",
        side_effect=RuntimeError("audit down"),
    )
    def test_audit_failure_does_not_break_mark(self, _mocked_create):
        _, _, resp = self._mark("a3@ut.edu.co", "EST-AUD3", "DOC-AUD3")
        self.assertEqual(resp.status_code, 201)
        self.assertFalse(AuditLog.objects.exists())


class AuditLogEndpointTests(TestCase):
    def setUp(self):
        self.admin = User.objects.create_user(
            email="auditadmin@ut.edu.co",
            google_sub="sub-auditadmin",
            role="ROLE_ADMIN",
        )
        self.professor, _, _ = _make_prof_group(email="auditprof@ut.edu.co")
        self.student, _ = _make_student(
            "auditstudent@ut.edu.co", code="EST-AUE", doc="DOC-AUE"
        )

    def _get(self, user):
        client = APIClient()
        client.force_login(user)
        return client.get("/api/attendance/audit-logs/")

    def test_admin_gets_200_paginated(self):
        resp = self._get(self.admin)
        self.assertEqual(resp.status_code, 200)
        self.assertIn("results", resp.data)

    def test_professor_is_403(self):
        self.assertEqual(self._get(self.professor).status_code, 403)

    def test_student_is_403(self):
        self.assertEqual(self._get(self.student).status_code, 403)

    def test_filters_by_event_type_and_date(self):
        from django.utils import timezone

        AuditLog.objects.create(event_type="login_success", resource="user:1")
        AuditLog.objects.create(event_type="logout", resource="user:1")
        client = APIClient()
        client.force_login(self.admin)
        resp = client.get(
            "/api/attendance/audit-logs/",
            {
                "event_type": "logout",
                "from": timezone.localdate().isoformat(),
            },
        )
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(
            [row["event_type"] for row in resp.data["results"]], ["logout"]
        )

    def test_invalid_date_is_400(self):
        client = APIClient()
        client.force_login(self.admin)
        resp = client.get(
            "/api/attendance/audit-logs/", {"from": "2026-13-01"}
        )
        self.assertEqual(resp.status_code, 400)


class ReportClassIdLimitTests(TestCase):
    def _client(self, email):
        user, _, group = _make_prof_group(email=email)
        client = APIClient()
        client.force_login(user)
        return client, group

    def test_more_than_50_class_ids_is_400(self):
        client, _ = self._client("replimit@ut.edu.co")
        ids = ",".join(str(uuid.uuid4()) for _ in range(51))
        resp = client.get("/api/attendance/reports/", {"class_id": ids})
        self.assertEqual(resp.status_code, 400)

    def test_duplicate_class_ids_do_not_duplicate_rows(self):
        client, group = self._client("repdup@ut.edu.co")
        session, _ = _make_session(_make_class(group))
        _, student = _make_student(
            "repdupest@ut.edu.co", code="EST-RD", doc="DOC-RD"
        )
        Attendance.objects.create(
            session=session,
            student=student,
            ip_address="127.0.0.1",
            user_agent="test",
        )
        class_id = str(session.scheduled_class_id)
        resp = client.get(
            "/api/attendance/reports/", {"class_id": f"{class_id},{class_id}"}
        )
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(len(resp.data["results"]), 1)


class ReportFormulaInjectionTests(TestCase):
    def test_xlsx_escapes_formula_like_student_name(self):
        user, _, group = _make_prof_group(email="xlsxprof@ut.edu.co")
        session, _ = _make_session(_make_class(group))
        _, student = _make_student("xlsxest@ut.edu.co", code="EST-XL", doc="DOC-XL")
        student.first_name = "=1+1"
        student.save(update_fields=["first_name"])
        Attendance.objects.create(
            session=session,
            student=student,
            ip_address="127.0.0.1",
            user_agent="test",
        )
        client = APIClient()
        client.force_login(user)
        resp = client.get("/api/attendance/reports/", {"format": "xlsx"})
        self.assertEqual(resp.status_code, 200)
        workbook = load_workbook(io.BytesIO(resp.content))
        value = workbook["Reporte"].cell(row=2, column=1).value
        self.assertTrue(str(value).startswith("'"))


class StudentProfileImmutabilityTests(TestCase):
    def setUp(self):
        from allauth.socialaccount.models import SocialAccount

        self.user = User.objects.create_user(
            email="immutable@ut.edu.co",
            google_sub="sub-immutable",
            role="ROLE_STUDENT",
        )
        SocialAccount.objects.create(
            user=self.user,
            provider="google",
            uid="sub-immutable",
            extra_data={"given_name": "Real", "family_name": "Estudiante"},
        )
        self.client = APIClient()
        self.client.force_login(self.user)
        self.body = {
            "student_code": "123456789012",
            "document_number": "123456",
            "phone_number": "3001234567",
            "address": "Calle 1 #2-3",
        }
        first = self.client.patch("/api/auth/me/", self.body, format="json")
        self.assertEqual(first.status_code, 200)

    def test_changed_student_code_is_400(self):
        resp = self.client.patch(
            "/api/auth/me/",
            {**self.body, "student_code": "999888777"},
            format="json",
        )
        self.assertEqual(resp.status_code, 400)
        self.assertIn("student_code", resp.data)
        student = Student.objects.get(user=self.user)
        self.assertEqual(student.student_code, "123456789012")

    def test_changed_document_number_is_400(self):
        resp = self.client.patch(
            "/api/auth/me/",
            {**self.body, "document_number": "999999"},
            format="json",
        )
        self.assertEqual(resp.status_code, 400)
        self.assertIn("document_number", resp.data)

    def test_same_codes_can_update_other_fields(self):
        resp = self.client.patch(
            "/api/auth/me/",
            {**self.body, "address": "Carrera 9 #10-11"},
            format="json",
        )
        self.assertEqual(resp.status_code, 200)
        student = Student.objects.get(user=self.user)
        self.assertEqual(student.address, "Carrera 9 #10-11")
        self.assertEqual(student.student_code, "123456789012")
