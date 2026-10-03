"""Módulo de reportes de asistencia (RF-PROF-09 / RF-ADM-04).

Todo vive en ``apps.attendance`` (no hay app nueva): el filtro, la
construcción del queryset y los exportadores XLSX/PDF. La vista está en
``views_reports.py`` y la ruta en ``urls.py`` (``reports/``).
"""

import html
import io
import uuid as uuid_mod
from datetime import datetime, time
from zoneinfo import ZoneInfo

from django.db.models import Q
from django.utils import timezone
from rest_framework import serializers
from rest_framework.exceptions import PermissionDenied

from apps.academic.models import ClassModality, ClassStatus
from apps.authentication.models import UserRole

BOGOTA_TZ = ZoneInfo("America/Bogota")

#: Tope de filas para descarga XLSX (el preview JSON va paginado).
MAX_EXPORT_ROWS = 5000
#: Tope menor para PDF: una sola Table en memoria (riesgo OOM con 5000).
#: El PDF legible ronda 400 filas; 1000 es el máximo operativo.
MAX_PDF_ROWS = 1000
#: Rango máximo permitido entre fecha inicial y final (días, inclusivo).
MAX_RANGE_DAYS = 366


class MultipleUUIDField(serializers.Field):
    """Acepta un UUID o una lista de UUIDs (``class_id`` multivalor).

    Tolera un valor único, lista/tupla y cadenas separadas por comas.
    """

    default_error_messages = {
        "invalid": "Identificador de clase inválido.",
    }

    def to_internal_value(self, data):
        if data is None or data == "":
            return []
        items = data if isinstance(data, (list, tuple)) else [data]
        expanded = []
        for item in items:
            if item is None:
                continue
            if isinstance(item, str) and "," in item:
                expanded.extend([p.strip() for p in item.split(",") if p.strip()])
            elif isinstance(item, str):
                if item.strip():
                    expanded.append(item.strip())
            else:
                expanded.append(item)
        result = []
        for value in expanded:
            try:
                result.append(uuid_mod.UUID(str(value)))
            except (ValueError, AttributeError, TypeError):
                self.fail("invalid")
        return result

    def to_representation(self, value):
        return [str(v) for v in value]


class HasLocationField(serializers.Field):
    """Normaliza ``has_location`` a ``true``/``false``/``all``."""

    default_error_messages = {
        "invalid": "Valor inválido para ubicación (true, false o all).",
    }

    _MAP = {
        "true": "true",
        "1": "true",
        "si": "true",
        "sí": "true",
        "yes": "true",
        "false": "false",
        "0": "false",
        "no": "false",
        "all": "all",
        "todos": "all",
        "todas": "all",
    }

    def to_internal_value(self, data):
        if data is None or data == "":
            return "all"
        if isinstance(data, bool):
            return "true" if data else "false"
        key = str(data).strip().lower()
        if key not in self._MAP:
            self.fail("invalid")
        return self._MAP[key]

    def to_representation(self, value):
        return value


class ReportFilterSerializer(serializers.Serializer):
    """Filtros del reporte. Todo opcional; errores 400 en español."""

    course_id = serializers.UUIDField(
        required=False,
        allow_null=True,
        error_messages={"invalid": "Identificador de curso inválido."},
    )
    group_id = serializers.UUIDField(
        required=False,
        allow_null=True,
        error_messages={"invalid": "Identificador de grupo inválido."},
    )
    class_id = MultipleUUIDField(required=False)
    session_id = serializers.UUIDField(
        required=False,
        allow_null=True,
        error_messages={"invalid": "Identificador de sesión inválido."},
    )
    date_from = serializers.DateField(
        required=False,
        input_formats=["%Y-%m-%d"],
        error_messages={"invalid": "Formato de fecha inválido, usa AAAA-MM-DD."},
    )
    date_to = serializers.DateField(
        required=False,
        input_formats=["%Y-%m-%d"],
        error_messages={"invalid": "Formato de fecha inválido, usa AAAA-MM-DD."},
    )
    registered_from = serializers.DateField(
        required=False,
        input_formats=["%Y-%m-%d"],
        error_messages={"invalid": "Formato de fecha inválido, usa AAAA-MM-DD."},
    )
    registered_to = serializers.DateField(
        required=False,
        input_formats=["%Y-%m-%d"],
        error_messages={"invalid": "Formato de fecha inválido, usa AAAA-MM-DD."},
    )
    document = serializers.CharField(required=False, allow_blank=True, max_length=50)
    student_code = serializers.CharField(required=False, allow_blank=True, max_length=50)
    search = serializers.CharField(required=False, allow_blank=True, max_length=100)
    class_status = serializers.ChoiceField(
        required=False,
        choices=[c for c, _ in ClassStatus.choices],
        error_messages={"invalid_choice": "Estado de clase inválido."},
    )
    modality = serializers.ChoiceField(
        required=False,
        choices=[c for c, _ in ClassModality.choices],
        error_messages={"invalid_choice": "Modalidad inválida."},
    )
    professor_id = serializers.UUIDField(
        required=False,
        allow_null=True,
        error_messages={"invalid": "Identificador de profesor inválido."},
    )
    has_location = HasLocationField(required=False)
    format = serializers.CharField(
        required=False,
        allow_blank=True,
        default="json",
        max_length=10,
    )

    def validate_format(self, value):
        norm = str(value or "json").strip().lower() or "json"
        if norm not in ("json", "xlsx", "pdf"):
            raise serializers.ValidationError(
                "Formato inválido (json, xlsx o pdf)."
            )
        return norm

    def validate(self, attrs):
        date_from = attrs.get("date_from")
        date_to = attrs.get("date_to")
        if date_from and date_to:
            if date_from > date_to:
                raise serializers.ValidationError(
                    {"date_from": "La fecha inicial no puede ser posterior a la final."}
                )
            # Inclusivo: (to - from).days + 1 son los días cubiertos.
            if (date_to - date_from).days + 1 > MAX_RANGE_DAYS:
                raise serializers.ValidationError(
                    {"date_to": "El rango de fechas no puede superar 366 días."}
                )
        reg_from = attrs.get("registered_from")
        reg_to = attrs.get("registered_to")
        if reg_from and reg_to:
            if reg_from > reg_to:
                raise serializers.ValidationError(
                    {
                        "registered_from": (
                            "La fecha inicial de registro no puede ser "
                            "posterior a la final."
                        )
                    }
                )
            if (reg_to - reg_from).days + 1 > MAX_RANGE_DAYS:
                raise serializers.ValidationError(
                    {
                        "registered_to": (
                            "El rango de fechas de registro no puede superar 366 días."
                        )
                    }
                )
        return attrs


def _day_bounds(day, end=False):
    """Convierte un ``date`` a aware en America/Bogota (inicio/fin del día)."""
    moment = datetime.combine(day, time.max if end else time.min)
    if timezone.is_naive(moment):
        moment = timezone.make_aware(moment, BOGOTA_TZ)
    return moment


def build_report_queryset(validated, user):
    """Construye el queryset del reporte según rol y filtros.

    - Profesor: solo marcaciones de sus grupos. Filtrar por grupo/clase/
      sesión ajenos responde 403; ``professor_id`` es solo admin (403).
    - ``course_id`` ajeno NO responde 403: los cursos son compartidos entre
      profesores, así que el filtro se aplica dentro de los grupos propios
      y un curso ajeno simplemente devuelve 200 vacío. Se documenta para
      no confundirlo con el 403 de grupo/clase/sesión.
    - Admin: todo, con filtro opcional por ``professor_id``.
    """
    from apps.academic.models import AcademicGroup, ScheduledClass
    from apps.attendance.models import Attendance, AttendanceSession
    from apps.authentication.models import Professor

    is_admin = getattr(user, "role", None) == UserRole.ROLE_ADMIN
    is_professor = getattr(user, "role", None) == UserRole.ROLE_PROFESSOR

    qs = Attendance.objects.all()

    professor_id = validated.get("professor_id")
    if professor_id and not is_admin:
        raise PermissionDenied("No tienes permiso para filtrar por profesor.")

    if is_professor:
        qs = qs.filter(session__scheduled_class__group__professor__user=user)
        group_id = validated.get("group_id")
        if group_id:
            group = AcademicGroup.objects.select_related("professor__user").filter(
                pk=group_id
            ).first()
            if group is not None and group.professor.user_id != user.id:
                raise PermissionDenied("No tienes acceso a ese grupo.")
        for class_pk in validated.get("class_id") or []:
            clase = ScheduledClass.objects.select_related("group__professor__user").filter(
                pk=class_pk
            ).first()
            if clase is not None and clase.group.professor.user_id != user.id:
                raise PermissionDenied("No tienes acceso a esa clase.")
        session_id = validated.get("session_id")
        if session_id:
            session = AttendanceSession.objects.select_related(
                "scheduled_class__group__professor__user"
            ).filter(pk=session_id).first()
            if (
                session is not None
                and session.scheduled_class.group.professor.user_id != user.id
            ):
                raise PermissionDenied("No tienes acceso a esa sesión.")
    elif is_admin and professor_id:
        if not Professor.objects.filter(pk=professor_id).exists():
            return Attendance.objects.none()
        qs = qs.filter(session__scheduled_class__group__professor_id=professor_id)

    if validated.get("course_id"):
        qs = qs.filter(session__scheduled_class__group__course_id=validated["course_id"])
    if validated.get("group_id"):
        qs = qs.filter(session__scheduled_class__group_id=validated["group_id"])
    if validated.get("class_id"):
        qs = qs.filter(session__scheduled_class_id__in=validated["class_id"])
    if validated.get("session_id"):
        qs = qs.filter(session_id=validated["session_id"])
    if validated.get("class_status"):
        qs = qs.filter(session__scheduled_class__status=validated["class_status"])
    if validated.get("modality"):
        qs = qs.filter(session__scheduled_class__modality=validated["modality"])

    document = (validated.get("document") or "").strip()
    if document:
        qs = qs.filter(student__document_number__icontains=document)
    student_code = (validated.get("student_code") or "").strip()
    if student_code:
        qs = qs.filter(student__student_code__icontains=student_code)
    search = (validated.get("search") or "").strip()
    if search:
        qs = qs.filter(
            Q(student__first_name__icontains=search)
            | Q(student__last_name__icontains=search)
            | Q(student__document_number__icontains=search)
            | Q(student__student_code__icontains=search)
            | Q(student__user__email__icontains=search)
        )

    if validated.get("date_from"):
        qs = qs.filter(
            session__scheduled_class__start_time__gte=_day_bounds(validated["date_from"])
        )
    if validated.get("date_to"):
        qs = qs.filter(
            session__scheduled_class__start_time__lte=_day_bounds(
                validated["date_to"], end=True
            )
        )
    if validated.get("registered_from"):
        qs = qs.filter(registered_at__gte=_day_bounds(validated["registered_from"]))
    if validated.get("registered_to"):
        qs = qs.filter(
            registered_at__lte=_day_bounds(validated["registered_to"], end=True)
        )

    has_location = validated.get("has_location") or "all"
    if has_location == "true":
        qs = qs.filter(latitude__isnull=False, longitude__isnull=False)
    elif has_location == "false":
        qs = qs.filter(Q(latitude__isnull=True) | Q(longitude__isnull=True))

    return qs.select_related(
        "student__user",
        "session__scheduled_class__group__course",
        "session__scheduled_class__group__professor__user",
    ).order_by("registered_at", "id")


def report_row(attendance):
    """Serializa una marcación a la fila JSON del reporte."""
    student = getattr(attendance, "student", None)
    session = getattr(attendance, "session", None)
    clase = getattr(session, "scheduled_class", None) if session else None
    group = getattr(clase, "group", None) if clase else None
    course = getattr(group, "course", None) if group else None
    email = getattr(getattr(student, "user", None), "email", None)
    full_name = ""
    if student is not None:
        full_name = f"{student.first_name} {student.last_name}".strip()
    return {
        "id": str(attendance.id),
        "student_name": full_name,
        "student_code": getattr(student, "student_code", None),
        "document_number": getattr(student, "document_number", None),
        "email": email,
        "course_code": getattr(course, "code", None),
        "course_name": getattr(course, "name", None),
        "group_id": str(group.id) if group else None,
        "group_code": getattr(group, "group_code", None),
        "term_period": getattr(group, "term_period", None),
        "class_id": str(clase.id) if clase else None,
        "class_title": getattr(clase, "title", None),
        "class_start_time": getattr(clase, "start_time", None),
        "modality": getattr(clase, "modality", None),
        "class_status": getattr(clase, "status", None),
        "session_id": str(session.id) if session else None,
        "registered_at": attendance.registered_at,
        "has_location": (
            attendance.latitude is not None and attendance.longitude is not None
        ),
    }


def _fmt_bogota(value):
    if not value:
        return ""
    return timezone.localtime(value, BOGOTA_TZ).strftime("%d/%m/%Y %H:%M")


_REPORT_COLUMNS = [
    "Estudiante",
    "Código",
    "Documento",
    "Correo",
    "Curso",
    "Grupo",
    "Clase",
    "Modalidad",
    "Estado",
    "Inicio clase",
    "Registrado",
    "Ubicación",
]

_REPORT_WIDTHS = [30, 12, 14, 28, 28, 10, 28, 12, 14, 17, 17, 11]


def _row_cells(attendance):
    row = report_row(attendance)
    group_label = ""
    if row["group_code"]:
        group_label = row["group_code"]
        if row["term_period"]:
            group_label = f"{group_label} ({row['term_period']})"
    return [
        row["student_name"] or "",
        row["student_code"] or "",
        row["document_number"] or "",
        row["email"] or "",
        row["course_name"] or "",
        group_label,
        row["class_title"] or "",
        row["modality"] or "",
        row["class_status"] or "",
        _fmt_bogota(row["class_start_time"]),
        _fmt_bogota(row["registered_at"]),
        "Sí" if row["has_location"] else "No",
    ]


def export_xlsx(queryset):
    """Genera el XLSX del reporte (cabecera #B3200E, fechas dd/mm/aaaa hh:mm)."""
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Reporte")

    header_fill = PatternFill("solid", fgColor="B3200E")
    header_font = Font(bold=True, color="FFFFFF")
    header_cells = []
    for title in _REPORT_COLUMNS:
        cell = WriteOnlyCell(ws, value=title)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center")
        header_cells.append(cell)
    ws.append(header_cells)

    for idx, width in enumerate(_REPORT_WIDTHS, start=1):
        try:
            ws.column_dimensions[get_column_letter(idx)].width = width
        except Exception:
            pass
    try:
        # WriteOnlyWorksheet ignora vistas; no fallar si no soporta freeze.
        ws.freeze_panes = "A2"
    except Exception:
        pass
    last_letter = get_column_letter(len(_REPORT_COLUMNS))
    try:
        total = queryset.count()
        ws.auto_filter.ref = f"A1:{last_letter}{total + 1}"
    except Exception:
        pass

    for attendance in queryset.iterator(chunk_size=1000):
        ws.append(_row_cells(attendance))

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _filters_summary(validated):
    parts = []
    labels = {
        "course_id": "Curso",
        "group_id": "Grupo",
        "session_id": "Sesión",
        "professor_id": "Profesor",
        "date_from": "Desde",
        "date_to": "Hasta",
        "registered_from": "Registro desde",
        "registered_to": "Registro hasta",
        "document": "Documento",
        "student_code": "Código",
        "search": "Búsqueda",
        "class_status": "Estado",
        "modality": "Modalidad",
        "has_location": "Ubicación",
    }
    if validated.get("class_id"):
        parts.append(f"Clases: {len(validated['class_id'])} seleccionada(s)")
    for key, label in labels.items():
        value = validated.get(key)
        if value not in (None, "", []):
            # Escapa input de usuario: Paragraph interpreta mini-HTML y
            # "<>&" sin escapar rompe el parse (500). Ver ALTO 1 QA.
            parts.append(f"{label}: {html.escape(str(value))}")
    return "; ".join(parts) if parts else "Sin filtros (todos los registros)"


def export_pdf(queryset, validated=None):
    """Genera el PDF del reporte (A4 horizontal, Helvetica, sin emojis)."""
    from reportlab.lib import colors
    from reportlab.lib.colors import HexColor
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    validated = validated or {}
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=12 * mm,
        rightMargin=12 * mm,
        topMargin=14 * mm,
        bottomMargin=14 * mm,
        title="Reporte de asistencia SIA-QR",
    )
    styles = getSampleStyleSheet()
    brand = HexColor("#B3200E")
    generated = timezone.localtime(timezone.now(), BOGOTA_TZ).strftime("%d/%m/%Y %H:%M")

    story = [
        Paragraph("Reporte de asistencia SIA-QR", styles["Heading1"]),
        Paragraph(f"Filtros: {_filters_summary(validated)}", styles["Normal"]),
        Paragraph(f"Generado: {generated} (America/Bogota)", styles["Normal"]),
        Spacer(1, 6),
    ]

    header = ["Estudiante", "Código", "Curso", "Clase", "Inicio", "Registrado", "Ubic."]
    data = [header]
    for attendance in queryset.iterator(chunk_size=1000):
        cells = _row_cells(attendance)
        data.append([cells[0], cells[1], cells[4], cells[6], cells[9], cells[10], cells[11]])

    usable = landscape(A4)[0] - 24 * mm
    col_widths = [
        usable * 0.24,
        usable * 0.10,
        usable * 0.20,
        usable * 0.20,
        usable * 0.10,
        usable * 0.10,
        usable * 0.06,
    ]
    table = Table(data, colWidths=col_widths, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), brand),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
            ]
        )
    )
    story.append(table)

    def _footer(canvas, _doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.drawRightString(
            landscape(A4)[0] - 12 * mm,
            10 * mm,
            f"Página {canvas.getPageNumber()}",
        )
        canvas.restoreState()

    doc.build(story, onFirstPage=_footer, onLaterPages=_footer)
    return buffer.getvalue()
