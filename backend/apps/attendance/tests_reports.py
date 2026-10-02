"""Tests del módulo de reportes (RF-PROF-09 / RF-ADM-04)."""

import hashlib
from datetime import timedelta
from unittest.mock import patch

from django.test import TestCase
from django.utils import timezone
from rest_framework.test import APIClient

from apps.academic.models import (
    AcademicGroup,
    ClassStatus,
    Course,
    ScheduledClass,
)
from apps.attendance.models import Attendance, AttendanceSession
from apps.authentication.models import (
    Professor,
    Student,
    User,
)

REPORT_URL = "/api/attendance/reports/"


def _rep_prof(email="reprof@ut.edu.co", code="EMP-REP"):
    user = User.objects.create_user(
        email=email, google_sub=f"sub-{email}", role="ROLE_PROFESSOR"
    )
    prof = Professor.objects.create(
        user=user, employee_code=code, first_name="Repo", last_name="Prof"
    )
    course = Course.objects.create(code=f"C-{code}", name="Curso Reporte")
    group = AcademicGroup.objects.create(
        course=course, professor=prof, group_code="01", term_period="2026-1"
    )
    return user, prof, group


def _rep_class(group, title="Clase Reporte", status=ClassStatus.COMPLETED):
    return ScheduledClass.objects.create(
        group=group,
        title=title,
        start_time=timezone.now() - timedelta(days=1),
        duration_minutes=60,
        qr_duration_minutes=10,
        modality="PRESENTIAL",
        status=status,
    )


def _rep_student(email, code, doc, first="Pena", last="Nino"):
    user = User.objects.create_user(
        email=email, google_sub=f"sub-{email}", role="ROLE_STUDENT"
    )
    student = Student.objects.create(
        user=user, student_code=code, document_number=doc,
        first_name=first, last_name=last,
        phone_number="+57 300 123 4567", address="Calle 1 #2-3",
    )
    return user, student


def _rep_session(clase, n=0):
    raw = f"rep-tok-{clase.id}-{n}-{AttendanceSession.objects.count()}"
    return AttendanceSession.objects.create(
        scheduled_class=clase,
        token_hash=hashlib.sha256(raw.encode()).hexdigest(),
        expires_at=timezone.now() + timedelta(minutes=10),
        is_active=False,
    )


def _rep_mark(session, student, lat=None, lon=None):
    return Attendance.objects.create(
        session=session, student=student,
        latitude=lat, longitude=lon, accuracy=None,
        ip_address="127.0.0.1", user_agent="qa",
    )


class AttendanceReportTests(TestCase):
    def setUp(self):
        self.owner, self.owner_prof, self.group = _rep_prof()
        self.clase = _rep_class(self.group)
        self.session = _rep_session(self.clase)
        _, self.student = _rep_student(
            "munoz@ut.edu.co", "EST-N1", "DOC-N1",
            first="María Peña", last="Muñoz Niño",
        )
        _rep_mark(self.session, self.student, lat="4.5", lon="-75.5")

        self.other, _, self.other_group = _rep_prof(
            email="otro@ut.edu.co", code="EMP-OTR"
        )
        self.other_class = _rep_class(self.other_group, title="Clase Ajena")
        self.other_session = _rep_session(self.other_class, n=99)
        _, self.other_student = _rep_student(
            "otro-est@ut.edu.co", "EST-N2", "DOC-N2"
        )
        _rep_mark(self.other_session, self.other_student)

        self.admin = User.objects.create_user(
            email="adm@ut.edu.co", google_sub="sub-adm", role="ROLE_ADMIN"
        )
        self.owner_client = APIClient()
        self.owner_client.force_login(self.owner)
        self.other_client = APIClient()
        self.other_client.force_login(self.other)
        self.admin_client = APIClient()
        self.admin_client.force_login(self.admin)

    def test_anonymous_is_401(self):
        resp = APIClient().get(REPORT_URL)
        self.assertEqual(resp.status_code, 401)

    def test_student_is_403(self):
        stu_user, _ = _rep_student("solo-est@ut.edu.co", "EST-X", "DOC-X")
        client = APIClient()
        client.force_login(stu_user)
        resp = client.get(REPORT_URL)
        self.assertEqual(resp.status_code, 403)

    def test_prof_sees_only_own(self):
        resp = self.owner_client.get(REPORT_URL)
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.data["count"], 1)
        row = resp.data["results"][0]
        self.assertIn("Muñoz", row["student_name"])
        self.assertTrue(row["has_location"])

    def test_prof_forbidden_other_group(self):
        resp = self.owner_client.get(
            REPORT_URL, {"group_id": str(self.other_group.id)}
        )
        self.assertEqual(resp.status_code, 403)

    def test_prof_forbidden_other_session(self):
        resp = self.owner_client.get(
            REPORT_URL, {"session_id": str(self.other_session.id)}
        )
        self.assertEqual(resp.status_code, 403)

    def test_prof_cannot_filter_professor_id(self):
        resp = self.owner_client.get(
            REPORT_URL, {"professor_id": str(self.owner_prof.id)}
        )
        self.assertEqual(resp.status_code, 403)

    def test_admin_filters_by_professor(self):
        resp = self.admin_client.get(
            REPORT_URL, {"professor_id": str(self.owner_prof.id)}
        )
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.data["count"], 1)
        resp_all = self.admin_client.get(REPORT_URL)
        self.assertEqual(resp_all.data["count"], 2)

    def test_search_filters(self):
        resp = self.admin_client.get(REPORT_URL, {"search": "muñoz"})
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.data["count"], 1)
        resp = self.admin_client.get(REPORT_URL, {"document": "DOC-N2"})
        self.assertEqual(resp.data["count"], 1)

    def test_has_location_filter(self):
        resp = self.admin_client.get(REPORT_URL, {"has_location": "false"})
        self.assertEqual(resp.data["count"], 1)
        resp = self.admin_client.get(REPORT_URL, {"has_location": "true"})
        self.assertEqual(resp.data["count"], 1)

    def test_invalid_range_is_400_spanish(self):
        resp = self.admin_client.get(
            REPORT_URL, {"date_from": "2026-01-01", "date_to": "2024-01-01"}
        )
        self.assertEqual(resp.status_code, 400)
        resp = self.admin_client.get(
            REPORT_URL, {"date_from": "2024-01-01", "date_to": "2026-01-01"}
        )
        self.assertEqual(resp.status_code, 400)
        self.assertIn("366", str(resp.data))

    def test_export_xlsx_headers_and_enie(self):
        resp = self.admin_client.get(REPORT_URL, {"format": "xlsx"})
        self.assertEqual(resp.status_code, 200)
        self.assertIn(
            "spreadsheetml.sheet", resp["Content-Type"]
        )
        self.assertIn("attachment", resp["Content-Disposition"])
        self.assertIn("filename*=UTF-8", resp["Content-Disposition"])
        self.assertIn("siaqr_reporte_", resp["Content-Disposition"])
        from openpyxl import load_workbook

        import io as _io

        wb = load_workbook(filename=_io.BytesIO(resp.content), read_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        self.assertEqual(rows[0][0], "Estudiante")
        self.assertTrue(any("Muñoz" in str(c) for r in rows[1:] for c in r))

    def test_export_pdf_headers(self):
        resp = self.admin_client.get(REPORT_URL, {"format": "pdf"})
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp["Content-Type"], "application/pdf")
        self.assertIn("attachment", resp["Content-Disposition"])
        self.assertIn("filename*=UTF-8", resp["Content-Disposition"])
        self.assertTrue(resp.content.startswith(b"%PDF"))

    def test_export_cap_5000(self):
        with patch(
            "apps.attendance.views_reports.MAX_EXPORT_ROWS", 1
        ):
            resp = self.admin_client.get(REPORT_URL, {"format": "xlsx"})
            self.assertEqual(resp.status_code, 400)
            self.assertIn("ximo", str(resp.data))
